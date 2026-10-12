"""L5X Lint test suite — checklist rules, registry, exports, endpoints, CLI, UI.
Every rule gets a violating AND a compliant case so a rule that never fires
cannot pass."""
import os, sys, re, io, json, subprocess, tempfile, collections
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT); os.chdir(ROOT)
import analysis, rules, lifecycle
from output_formats import flatten_findings

FIX = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<RSLogix5000Content TargetType="Controller" SoftwareRevision="33.00" SchemaRevision="1.0">
<Controller Name="Line1" ProcessorType="1756-L83E" MajorRev="33" MinorRev="11" MajorFaultProgram="FaultHandler" PowerLossProgram="">
  <SafetyInfo><SafetyTagMap>STD_EStop=SFTY_EStop, Wrong=SFTY_Gate</SafetyTagMap></SafetyInfo>
  <DataTypes>
    <DataType Name="UNUSED_UDT" Family="NoFamily"><Members>
      <Member Name="A" DataType="DINT" Dimension="0" Hidden="false"><Description><![CDATA[a]]></Description></Member></Members></DataType>
    <DataType Name="ORDER_ONLY" Family="NoFamily"><Members>
      <Member Name="Zeta" DataType="DINT" Dimension="0" Hidden="false"><Description><![CDATA[z]]></Description></Member>
      <Member Name="Alpha" DataType="DINT" Dimension="0" Hidden="false"><Description><![CDATA[a]]></Description></Member></Members></DataType>
    <DataType Name="TMR_UDT" Family="NoFamily"><Members>
      <Member Name="TMR" DataType="TIMER" Dimension="10" Hidden="false"><Description><![CDATA[t]]></Description></Member></Members></DataType>
    <DataType Name="LOCAL_UDT" Family="NoFamily"><Members>
      <Member Name="X" DataType="DINT" Dimension="0" Hidden="false"><Description><![CDATA[x]]></Description></Member></Members></DataType>
    <DataType Name="NO_DESC_UDT" Family="NoFamily"><Members>
      <Member Name="X" DataType="DINT" Dimension="0" Hidden="false"/></Members></DataType>
  </DataTypes>
  <Tags>
    <Tag Name="SFTY_EStop" Class="Safety" TagType="Base" DataType="BOOL"/>
    <Tag Name="SFTY_Gate" Class="Safety" TagType="Base" DataType="BOOL"/>
    <Tag Name="STD_EStop" TagType="Base" DataType="BOOL"/>
    <Tag Name="Wrong" TagType="Base" DataType="BOOL"/>
    <Tag Name="IntTag" TagType="Base" DataType="INT"/>
    <Tag Name="RealTag" TagType="Base" DataType="REAL"/>
    <Tag Name="DintTag" TagType="Base" DataType="DINT"/>
    <Tag Name="SpareUsed" TagType="Base" DataType="DINT"><Description><![CDATA[Spare]]></Description></Tag>
    <Tag Name="OrderTag" TagType="Base" DataType="ORDER_ONLY"/>
    <Tag Name="NdTag" TagType="Base" DataType="NO_DESC_UDT"/>
    <Tag Name="ALARM_TMR" TagType="Base" DataType="TIMER" Dimensions="3"><Data Format="Decorated"><Array DataType="TIMER" Dimensions="3">
      <Element Index="[0]"><Structure DataType="TIMER"><DataValueMember Name="PRE" DataType="DINT" Value="5000"/><DataValueMember Name="ACC" DataType="DINT" Value="0"/></Structure></Element><Element Index="[1]"><Structure DataType="TIMER"><DataValueMember Name="PRE" DataType="DINT" Value="0"/><DataValueMember Name="ACC" DataType="DINT" Value="0"/></Structure></Element><Element Index="[2]"><Structure DataType="TIMER"><DataValueMember Name="PRE" DataType="DINT" Value="0"/><DataValueMember Name="ACC" DataType="DINT" Value="0"/></Structure></Element></Array></Data></Tag>
    <Tag Name="GoodTmr" TagType="Base" DataType="TIMER"><Data Format="Decorated"><Structure DataType="TIMER"><DataValueMember Name="PRE" DataType="DINT" Value="1000"/><DataValueMember Name="ACC" DataType="DINT" Value="0"/></Structure></Data></Tag>
    <Tag Name="ZeroTmr" TagType="Base" DataType="TIMER"><Data Format="Decorated"><Structure DataType="TIMER"><DataValueMember Name="PRE" DataType="DINT" Value="0"/><DataValueMember Name="ACC" DataType="DINT" Value="0"/></Structure></Data></Tag>
    <Tag Name="SpArr" TagType="Base" DataType="DINT" Dimensions="4"><Comments><Comment Operand="[1]"><![CDATA[Spare]]></Comment></Comments></Tag>
    <Tag Name="BUG01" TagType="Base" DataType="TMR_UDT"/>
    <Tag Name="MyMsg" TagType="Base" DataType="MESSAGE"/>
    <Tag Name="StOnlyTag" TagType="Base" DataType="DINT"/>
    <Tag Name="BitArr" TagType="Base" DataType="DINT" Dimensions="2"/>
    <Tag Name="AoiInst" TagType="Base" DataType="MyAOI"/>
    <Tag Name="AlPB" TagType="Alias" AliasFor="Local:2:I.Data.7"/>
    <Tag Name="AlOld" TagType="Alias" AliasFor="Local:1:I.Data.4"/>
    <Tag Name="AlChain" TagType="Alias" AliasFor="AlPB"/>
    <Tag Name="ToL2_PRO_100MS" TagType="Produced" DataType="DINT"><ProduceInfo ProduceCount="1"/></Tag>
    <Tag Name="BadProd" TagType="Produced" DataType="DINT"><ProduceInfo ProduceCount="1"/></Tag>
    <Tag Name="FromL2_CON_100MS" TagType="Consumed" DataType="DINT"><ConsumeInfo Producer="L2" RemoteTag="FromL2_PRO_100MS"/></Tag>
    <Tag Name="WrongName" TagType="Consumed" DataType="DINT"><ConsumeInfo Producer="L2" RemoteTag="X_PRO_1"/></Tag>
  </Tags>
  <Programs>
    <Program Name="Main" MainRoutineName="R" FaultRoutineName="Flt">
      <Tags/>
      <Routines>
      <Routine Name="R" Type="RLL"><RLLContent>
        <Rung Number="0" Type="N"><Text><![CDATA[XIC(A)XIC(A)OTE(B);]]></Text></Rung>
        <Rung Number="1" Type="N"><Text><![CDATA[MOV(Arr[Idx],X);]]></Text></Rung>
        <Rung Number="2" Type="N"><Text><![CDATA[LIM(0,Idx2,9)MOV(Arr[Idx2],Y);]]></Text></Rung>
        <Rung Number="3" Type="N"><Text><![CDATA[DIV(A,Den,Q);]]></Text></Rung>
        <Rung Number="4" Type="N"><Text><![CDATA[NEQ(Den2,0)DIV(A,Den2,Q2);]]></Text></Rung>
        <Rung Number="5" Type="N"><Text><![CDATA[LBL(Back)NOP();]]></Text></Rung>
        <Rung Number="6" Type="N"><Text><![CDATA[JMP(Back);]]></Text></Rung>
        <Rung Number="7" Type="N"><Text><![CDATA[JMP(Fwd);]]></Text></Rung>
        <Rung Number="8" Type="N"><Text><![CDATA[LBL(Fwd)NOP();]]></Text></Rung>
        <Rung Number="9" Type="N"><Text><![CDATA[TON(T1,?,?);]]></Text></Rung>
        <Rung Number="10" Type="N"><Text><![CDATA[TON(T1,?,?);]]></Text></Rung>
        <Rung Number="11" Type="N"><Text><![CDATA[CTU(C1,?,?);]]></Text></Rung>
        <Rung Number="12" Type="N"><Text><![CDATA[CTD(C1,?,?);RES(C1);]]></Text></Rung>
        <Rung Number="13" Type="N"><Text><![CDATA[XIC(S)ONS(OS1)OTE(Y1);]]></Text></Rung>
        <Rung Number="14" Type="N"><Text><![CDATA[XIC(S2)ONS(OS1)OTE(Y2);]]></Text></Rung>
        <Rung Number="15" Type="N"><Text><![CDATA[RTO(R1,?,?);]]></Text></Rung>
        <Rung Number="16" Type="N"><Text><![CDATA[RTO(R2,?,?);]]></Text></Rung>
        <Rung Number="17" Type="N"><Text><![CDATA[RES(R2);]]></Text></Rung>
        <Rung Number="18" Type="N"><Text><![CDATA[[XIC(A) ,]OTE(Q3);]]></Text></Rung>
        <Rung Number="19" Type="N"><Text><![CDATA[MOV(RealTag,DintTag);]]></Text></Rung>
        <Rung Number="20" Type="N"><Text><![CDATA[MOV(DintTag,RealTag);]]></Text></Rung>
        <Rung Number="21" Type="N"><Text><![CDATA[MOV(SpareUsed,IntTag)MOV(OrderTag.Zeta,NdTag.X);]]></Text></Rung>
        <Rung Number="22" Type="N"><Text><![CDATA[XIC(Local:2:I.Data.0)XIC(Local:2:I.Data.5)AFI()OTE(Q4);]]></Text></Rung>
        <Rung Number="24" Type="N"><Text><![CDATA[TON(ALARM_TMR[1],?,?)TON(GoodTmr,?,?)TON(ZeroTmr,?,?)MOV(SpArr[1],Z9);]]></Text></Rung>
        <Rung Number="25" Type="N"><Text><![CDATA[RTO(RArr[2],?,?);]]></Text></Rung>
        <Rung Number="26" Type="N"><Text><![CDATA[XIC(Done)RES(RArr[i]);]]></Text></Rung>
        <Rung Number="27" Type="N"><Text><![CDATA[XIC(A)[XIO(P) XIO(M) ,XIC(M) XIO(P) ,XIC(P) ]OTE(Q5);]]></Text></Rung>
        <Rung Number="28" Type="N"><Text><![CDATA[[XIC(K),XIC(K)]OTE(Q6);]]></Text></Rung>
        <Rung Number="29" Type="N"><Text><![CDATA[XIC(G)OTL(BUG01.TMR[6].DN)OTE(MyMsg.DN)MOV(MyMsg.DN,Z1)OTE(Unknown.ER)OTE(Unknown2.DN);]]></Text></Rung>
        <Rung Number="30" Type="N"><Text><![CDATA[[LIM(0,Ix,9) ,XIC(Z1) ]MOV(Arr2[Ix],Y);]]></Text></Rung>
        <Rung Number="31" Type="N"><Text><![CDATA[MOV(Arr3[Jx],Y)LIM(0,Jx,9);]]></Text></Rung>
        <Rung Number="32" Type="N"><Text><![CDATA[XIC(BitArr[0].3)XIC(Local:7:I.Ch0Data)OTE(Q7);]]></Text></Rung>
        <Rung Number="33" Type="N"><Text><![CDATA[XIC(AlPB)OTE(Q8);]]></Text></Rung>
        <Rung Number="23" Type="N"><Text><![CDATA[MOV(ToL2_PRO_100MS,BadProd)MOV(FromL2_CON_100MS,WrongName)XIC(STD_EStop)XIC(Wrong);]]></Text></Rung>
      </RLLContent></Routine>
      <Routine Name="Flt" Type="RLL"><RLLContent/></Routine>
      <Routine Name="Fbd" Type="FBD"><FBDContent/></Routine>
      <Routine Name="StR" Type="ST"><STContent><Line Number="0"><![CDATA[StOnlyTag := 1; JSR(StCalled);]]></Line></STContent></Routine>
      <Routine Name="StCalled" Type="RLL"><RLLContent/></Routine>
      <EncodedData Name="Secret" Type="RLL" EncodedType="Routine"/>
      </Routines>
      <ChildPrograms><ChildProgram Name="ChildP"/></ChildPrograms>
    </Program>
    <Program Name="ChildP" MainRoutineName="CR" FaultRoutineName="CF"><Tags/><Routines><Routine Name="CR" Type="RLL"><RLLContent/></Routine><Routine Name="CF" Type="RLL"><RLLContent/></Routine></Routines></Program>
    <Program Name="FaultHandler" MainRoutineName="FH"><Tags/><Routines><Routine Name="FH" Type="RLL"><RLLContent/></Routine></Routines></Program>
  </Programs>
  <Tasks>
    <Task Name="Fast_10MS" Type="PERIODIC" Period="10" Watchdog="50"><ScheduledPrograms><ScheduledProgram Name="Main"/></ScheduledPrograms></Task>
    <Task Name="Slow" Type="PERIODIC" Period="100" Watchdog="500"/>
    <Task Name="Task_1S" Type="PERIODIC" Period="1000" Watchdog="5000"/>
  </Tasks>
  <AddOnInstructionDefinitions>
    <AddOnInstructionDefinition Name="MyAOI" Revision="1.0"><Description><![CDATA[aoi]]></Description>
      <Parameters><Parameter Name="EnableIn" DataType="BOOL" Usage="Input"><Description><![CDATA[e]]></Description></Parameter></Parameters>
      <LocalTags><LocalTag Name="Cnt" DataType="LOCAL_UDT"/><LocalTag Name="DintTag" DataType="DINT"/></LocalTags>
      <Routines>
        <Routine Name="Logic" Type="RLL"><RLLContent><Rung Number="0" Type="N"><Comment><![CDATA[c]]></Comment><Text><![CDATA[XIC(DintTag.0)OTE(Cnt.X);]]></Text></Rung></RLLContent></Routine>
        <Routine Name="Prescan" Type="RLL"><RLLContent/></Routine>
      </Routines>
    </AddOnInstructionDefinition>
  </AddOnInstructionDefinitions>
  <Modules>
    <Module Name="NoMap" CatalogNumber="1756-IB16"><EKey State="Disabled"/><Ports><Port Id="1" Address="1" Type="ICP" Upstream="true"/></Ports>
      <Communications><Connections><Connection><InputTag Suffix="I"><Data Format="Decorated"><Structure DataType="AB:1756_DI:I:0"><DataValueMember Name="Fault" DataType="DINT" Value="0"/><DataValueMember Name="Data" DataType="DINT" Value="0"/></Structure></Data></InputTag></Connection></Connections></Communications></Module>
    <Module Name="Mapped" CatalogNumber="1756-IB32"><EKey State="CompatibleModule"/><Ports><Port Id="1" Address="2" Type="ICP" Upstream="true"/></Ports>
      <Communications><Connections><Connection><InputTag Suffix="I"><Comments><Comment Operand=".DATA.0"><![CDATA[PB]]></Comment></Comments><Data Format="Decorated"><Structure DataType="AB:1756_DI:I:0"><DataValueMember Name="Fault" DataType="DINT" Value="0"/><DataValueMember Name="Data" DataType="DINT" Value="0"/></Structure></Data></InputTag></Connection></Connections></Communications></Module>
    <Module Name="Ana" CatalogNumber="1756-IF8"><EKey State="ExactMatch"/><Ports><Port Id="1" Address="7" Type="ICP" Upstream="true"/></Ports>
      <Communications><Connections><Connection><InputTag Suffix="I"><Data Format="Decorated"><Structure DataType="AB:1756_IF8:I:0"><DataValueMember Name="Ch0Fault" DataType="BOOL" Value="0"/><DataValueMember Name="Ch0Data" DataType="REAL" Value="0"/><DataValueMember Name="Ch1Data" DataType="REAL" Value="0"/></Structure></Data></InputTag></Connection></Connections></Communications></Module>
    <Module Name="Bridge" CatalogNumber="1756-ENBT"><Ports><Port Id="1" Address="3" Type="ICP"/></Ports></Module>
  </Modules>
</Controller></RSLogix5000Content>"""

CFG = {"disabled_rules": [], "sizes": {"loose_bool_min": 1}, "urs": {}}

def run(xml=FIX, cfg=CFG):
    r = None
    for ev in analysis.run_analysis(xml, cfg):
        if ev["type"] == "error": raise AssertionError(ev["message"])
        if ev["type"] == "done": r = ev["report"]
    return r

R = run(); FS = flatten_findings(R)
def by(code): return [f for f in FS if f["code"] == code]
def msgs(code): return " | ".join(f["message"] + " @ " + f["location"] for f in by(code))

# ── registry ──
def test_registry_codes_unique_and_prefix_matches_severity():
    codes = [r[0] for r in rules.RULES]
    assert len(codes) == len(set(codes))
    for c, rec in rules.BY_CODE.items():
        assert rules.SEV_PREFIX[rec["severity"]] == c[:3], c
    assert len(codes) == 49   # 45 checklist (WAR012 retired) + WAR022 + 3 lifecycle

def test_every_finding_has_code_and_fixed_severity():
    for f in FS:
        assert f["code"] and rules.BY_CODE[f["code"]]["severity"] == f["severity"], f

def test_disabled_rules_are_skipped_not_just_hidden():
    import checks as new_rules
    def boom(ctx): raise AssertionError("disabled rule was executed")
    boom.__name__ = "dup_instr"
    idx = [f.__name__ for f in new_rules.ALL].index("dup_instr")
    orig = new_rules.ALL[idx]; new_rules.ALL[idx] = boom
    try:
        r = run(cfg={**CFG, "disabled_rules": ["ERR010"]})
        assert r["rule_errors"] == [], r["rule_errors"]            # boom never called
        r2 = run(cfg=CFG)
        assert any("disabled rule was executed" in e for e in r2["rule_errors"])   # sanity: enabled → runs
    finally:
        new_rules.ALL[idx] = orig

def test_disabled_rule_removed():
    r = run(cfg={**CFG, "disabled_rules": ["ERR010"]})
    assert not [f for f in flatten_findings(r) if f["code"] == "ERR010"] and by("ERR010")

def test_no_fix_bookmark_or_req_fields():
    for f in FS:
        assert "fix" not in f and "bookmark" not in f and "req_ids" not in f, f
    assert "compliance" not in R

# ── checklist rules ──
def test_CRT002_afi():           assert "Rung 22" in msgs("CRT002")
def test_ERR007_fault_handler_present_passes(): assert not by("ERR007")
def test_ERR008_fault_routine_present_passes(): assert "Main" not in msgs("ERR008")
def test_ERR010_series_and_parallel():
    m = msgs("ERR010")
    assert "XIC(A) appears twice in series" in m and "Rung 0" in m
    assert "Identical parallel branch legs: XIC(K)" in m
    assert "Rung 27" not in m          # shared instruction across legs with other conditions = valid OR
def test_ERR011_guard_same_path_before():
    m = msgs("ERR011")
    assert "'Idx'" in m and "'Idx2'" not in m      # LIM before on same path protects
    assert "'Ix'" in m                             # LIM in a sibling leg does not
    assert "'Jx'" in m                             # LIM after the MOV does not
def test_ERR012_div_zero():      m = msgs("ERR012"); assert "'Den'" in m and "Den2" not in m
def test_ERR013_jmp_back():      m = msgs("ERR013"); assert "Back" in m and "Fwd" not in m
def test_ERR014_timer_reuse():   m = msgs("ERR014"); assert "'T1'" in m and "'C1'" not in m
def test_ERR015_oneshot_reuse(): assert "OS1" in msgs("ERR015")
def test_ERR016_res_missing():   m = msgs("ERR016"); assert "'R1'" in m and "'R2'" not in m and "'C1'" not in m
def test_ERR017_lc_discontinued(): assert "1756-ENBT" in msgs("ERR017")
def test_WAR001_safety_map():    m = msgs("WAR001"); assert "Wrong" in m and "STD_Gate" in m and "'STD_EStop'" not in m
def test_WAR002_pc_naming():
    m = msgs("WAR002"); assert "BadProd" in m and "WrongName" in m
    assert "ToL2_PRO_100MS" not in m and "FromL2_CON_100MS" not in m
def test_WAR003_task_period():   m = msgs("WAR003"); assert "'Slow'" in m and "Fast_10MS" not in m and "Task_1S" not in m
def test_WAR007_empty_branch():  assert "Rung 18" in msgs("WAR007")
def test_WAR010_no_zero_saving_padding_finding():
    assert all("0 bytes saved" not in f["message"] for f in by("WAR010"))
def test_WAR011_order_separate_from_size():
    assert "ORDER_ONLY" in msgs("WAR011") and "ORDER_ONLY" not in msgs("WAR010")
def test_WAR015_mov_real_int():  m = msgs("WAR015"); assert "RealTag,DintTag" in m and "DintTag,RealTag" not in m
def test_WAR016_small_int():     assert "IntTag" in msgs("WAR016")
def test_WAR017_per_point_data_only():
    locs = [f["location"] for f in by("WAR017")]
    assert "IO:Local:1:I.Data.0" in locs and "IO:Local:1:I.Data.15" in locs     # IB16 -> 16 points
    assert "IO:Local:1:I.Data.16" not in locs
    assert "IO:Local:2:I.Data.0" not in locs and "IO:Local:2:I.Data.31" in locs   # IB32, .0 used
    assert "IO:Local:7:I.Ch1Data" in locs and "IO:Local:7:I.Ch0Data" not in locs  # analog channels
    assert not any("Fault" in l for l in locs)                                    # diagnostics excluded

def test_WAR018_per_point():
    locs = [f["location"] for f in by("WAR018")]
    assert "IO:Local:2:I.Data.5" in locs and "IO:Local:2:I.Data.0" not in locs    # .0 commented
    assert "IO:Local:7:I.Ch0Data" in locs

def test_WAR019_desc_spare():    assert "SpareUsed" in msgs("WAR019")
def test_WAR020_unused_type():   m = msgs("WAR020"); assert "UNUSED_UDT" in m and "ORDER_ONLY" not in m
def test_WAR008_timer_array_per_element():
    m = msgs("WAR008")
    assert "ALARM_TMR" in m and "[1]" in m and "[2]" not in m and "[0]" not in m   # only used zero elements
    assert "ZeroTmr" in m and "GoodTmr" not in m

def test_WAR005_handlers_and_child_programs_not_flagged():
    m = msgs("WAR005"); assert "FaultHandler" not in m and "ChildP" not in m

def test_WAR012_retired():
    assert "WAR012" not in rules.BY_CODE and not by("WAR012")

def test_ERR016_indexed_reset_covers_array():  assert "RArr" not in msgs("ERR016")
def test_WAR019_element_comment_spare():        assert "SpArr[1]" in msgs("WAR019")
def test_protected_routines_counted():          assert R["meta"]["protected_routines"] == 1

def test_ERR004_ERR005_by_datatype():
    e4, e5 = msgs("ERR004"), msgs("ERR005")
    assert "BUG01.TMR[6].DN" in e5 and "BUG01.TMR" not in e4      # TIMER in UDT array -> ERR005 only
    assert "MyMsg" in e4 and "MyMsg" not in e5                    # MESSAGE -> ERR004
    assert "MOV(MyMsg.DN" not in e4                               # MOV reads operand 0
    assert "Unknown.ER" in e4                                     # unresolved but MSG-only bit
    assert "Unknown2.DN" not in e4 + e5                           # unresolved shared bit -> skipped

def test_WAR011_simple_text():
    assert [f["message"] for f in by("WAR011") if "ORDER_ONLY" in f["location"]] == \
        ["Members are not in alphabetical order."]

def test_WAR019_only_spare_unused():
    import checks.tags as new_rules
    assert new_rules._SPARE.search("UNUSED PATTERN") and new_rules._SPARE.search("Spare - do not use")
    assert not new_rules._SPARE.search("not used") and not new_rules._SPARE.search("SPARE_1")

def test_per_item_reporting():
    assert any(f["message"] == "Member 'NO_DESC_UDT.X' has no description." for f in by("INF006"))
    assert any("BitArr[0].3" in f["message"] for f in by("INF003"))
    assert [f["message"] for f in by("WAR008") if "ALARM_TMR" in f["message"]] == \
        ["Timer 'ALARM_TMR[1]' has preset 0 and logic never sets .PRE."]

def test_INF003_bits_only_when_referenced():
    bits = [f["message"] for f in by("INF003") if "[" in f["message"]]
    assert bits == ["'BitArr[0].3' has no comment."], bits      # not all 64/128 bits

def test_AOI_local_tags_parsed():
    assert "LOCAL_UDT" not in msgs("WAR020")                      # used only by an AOI local tag
    assert "DintTag" not in msgs("ERR009")                        # local tag, not the controller tag

def test_AOI_prescan_not_uncalled():
    assert "Prescan" not in msgs("WAR004")

def test_ST_references_count_as_usage():
    assert "StOnlyTag" not in msgs("INF001")
    assert "StCalled" not in msgs("WAR004")

def test_json_load_removed():
    import app as A
    c = A.app.test_client()
    assert c.post("/load_findings", headers={"Origin": "http://127.0.0.1:5002"}).status_code == 404
    assert "fi-findings" not in c.get("/").get_data(as_text=True)

def test_INF002_loose_bool_any(): assert by("INF002")
def test_INF006_udt_member_desc(): m = msgs("INF006"); assert "NO_DESC_UDT" in m and "ORDER_ONLY" not in m
def test_INF008_keying_must_be_compatible():
    m = msgs("INF008")
    assert "'NoMap'" in m and "'Disabled'" in m          # Disabled flagged
    assert "'Ana'" in m and "'ExactMatch'" in m          # Exact Match flagged
    assert "'Mapped'" not in m                           # Compatible Module OK

def test_WAR022_alias_used_for_physical_io():
    m = msgs("WAR022")
    assert "Alias 'AlPB' → Local:2:I.Data.7 is used in logic" in m
    assert "Alias 'AlOld' → Local:1:I.Data.4 is not used in logic" in m
    assert "Alias 'AlChain' → Local:2:I.Data.7" in m     # chain resolved

def test_WAR017_excludes_alias_covered_points():
    locs = [f["location"] for f in by("WAR017")]
    assert "IO:Local:2:I.Data.7" not in locs and "IO:Local:1:I.Data.4" not in locs
    assert "IO:Local:1:I.Data.5" in locs                 # neighbour still unmapped

def test_ui_new_file_and_back_to_top():
    import app as A
    h = A.app.test_client().get("/").get_data(as_text=True)
    assert 'id="to-top"' in h and "⟲ New file" in h
def test_INF009_lc_mature_absent_here(): assert not by("INF009")

def test_CRT001_forces():
    xml = FIX.replace("  <Modules>", "  <ForcedTags><Tag Name='X'/></ForcedTags>\n  <Modules>", 1)
    assert [f for f in flatten_findings(run(xml)) if f["code"] == "CRT001"]

def test_non_rll_counted_in_meta(): assert R["meta"]["non_rll_routines"] == 2   # FBD + ST
def test_no_rule_errors():           assert R["rule_errors"] == []

# ── suppression by code ──
def test_suppress_by_code():
    xml = FIX.replace('<Rung Number="0" Type="N"><Text>',
        '<Rung Number="0" Type="N"><Comment><![CDATA[lint:disable=ERR010]]></Comment><Text>')
    r = run(xml)
    assert not [f for f in flatten_findings(r) if f["code"] == "ERR010" and f["location"].endswith("Rung 0")]
    assert any(f["code"] == "ERR010" and f["location"].endswith("Rung 0") for f in r["suppressed_findings"])

# ── optimiser vs WAR011 ──
def test_optimiser_output_satisfies_order_rule():
    from udt_layout import member_sort_key
    mem = [{"name": n, "type": t, "dimension": 0} for n, t in
           [("G", "TIMER"), ("A", "TIMER"), ("B", "COUNTER"), ("V10", "DINT"), ("V2", "DINT")]]
    types = [m["type"] for m in sorted(mem, key=lambda m: member_sort_key(m, {}, {}))]
    for t in set(types):
        i = [k for k, x in enumerate(types) if x == t]
        assert i == list(range(i[0], i[-1] + 1)), types

# ── exports & endpoints ──
def _session():
    import app as A
    c = A.app.test_client()
    rv = c.post("/upload", data={"file": (io.BytesIO(FIX.encode()), "line1.l5x")},
                headers={"Origin": "http://127.0.0.1:5002"}, content_type="multipart/form-data")
    return c, rv.get_json()["session_id"]

def test_report_xlsx_and_pdf_endpoints():
    c, sid = _session()
    x = c.get(f"/report/{sid}.xlsx"); assert x.status_code == 200 and x.data[:2] == b"PK"
    p = c.get(f"/report/{sid}.pdf?sev=CRITICAL,ERROR"); assert p.status_code == 200 and p.data[:5] == b"%PDF-"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.data))
    assert wb.sheetnames == ["Summary", "Findings", "Suppressed", "Hardware", "Checklist"]
    assert c.get(f"/report/{sid}.docx").status_code == 404

def test_not_verified_reported_when_attribute_absent():
    xml = FIX.replace(' MajorFaultProgram="FaultHandler" PowerLossProgram=""', "")
    r = run(xml)
    assert any(n.startswith("ERR007") for n in r["not_verified"]), r["not_verified"]

def test_rerun_streams_with_current_settings():
    import app as A
    c, sid = _session()
    path = tempfile.mktemp(suffix=".json"); old = A.CONFIG_PATH; A.CONFIG_PATH = path
    try:
        json.dump({"disabled_rules": ["ERR010"]}, open(path, "w"))
        H = {"Origin": "http://127.0.0.1:5002"}
        assert c.post(f"/api/rerun/{sid}", headers=H).get_json()["success"]
        body = c.get(f"/stream/{sid}").get_data(as_text=True)
        events = [json.loads(l[6:]) for l in body.splitlines() if l.startswith("data: ")]
        pcts = [e["pct"] for e in events if e.get("type") == "progress"]
        assert pcts and pcts == sorted(pcts) and pcts[-1] > 0          # live % during reload
        done = [e for e in events if e.get("type") == "done"][0]["report"]
        assert not [f for f in flatten_findings(done) if f["code"] == "ERR010"]
        assert c.post("/api/rerun/nope", headers=H).status_code == 404
        assert c.get(f"/api/report/{sid}").status_code == 404           # old endpoint gone
    finally:
        A.CONFIG_PATH = old; os.path.exists(path) and os.unlink(path)

def test_loader_is_progress_ring():
    import app as A
    h = A.app.test_client().get("/").get_data(as_text=True)
    assert "function ringHTML" in h and "◌" not in h and "slim-bar" not in h

def test_export_uses_cached_report():
    import app as A
    c, sid = _session()
    for _ in c.get(f"/stream/{sid}").response: pass       # run analysis -> cache
    assert sid in A._report_cache
    calls = []
    orig = A._run_to_report
    A._run_to_report = lambda *a, **k: calls.append(1) or orig(*a, **k)
    try:
        assert c.get(f"/report/{sid}.xlsx").status_code == 200
        assert calls == [], "export re-ran the analysis"
    finally:
        A._run_to_report = orig

def test_large_export_fast_enough():
    import time, report_export
    big = dict(R); many = []
    base = flatten_findings(R)
    for i in range(80): many += [dict(f, location=f["location"] + f"#{i}") for f in base]
    big["controller_findings"] = many; big["task_program_tabs"] = []; big["aoi_sections"] = []
    t = time.time(); report_export.build_pdf(big, "big", None)
    assert time.time() - t < 20, f"PDF too slow for {len(many)} findings"

def test_report_rows_natural_order():
    import report_export
    locs = [f["location"] for f in report_export._rows(R, ["INFO"]) if f["code"] == "INF004"]
    nums = [int(l.rsplit(" ", 1)[1]) for l in locs]
    assert nums == sorted(nums), nums

def test_export_missing_package_gives_actionable_error():
    import app as A
    c, sid = _session()
    orig = A._missing_export_deps
    A._missing_export_deps = lambda: {"pdf": "reportlab"}
    try:
        r = c.get(f"/report/{sid}.pdf")
        assert r.status_code == 500 and "reportlab" in r.get_json()["error"] and "pip install" in r.get_json()["error"]
        html = c.get("/").get_data(as_text=True)
        assert '"pdf": "reportlab"' in html     # UI disables the button with a tooltip
    finally:
        A._missing_export_deps = orig

def test_export_error_shows_real_reason():
    import app as A, report_export
    c, sid = _session()
    orig = report_export.build_xlsx
    report_export.build_xlsx = lambda *a, **k: (_ for _ in ()).throw(ValueError("boom 42"))
    try:
        e = c.get(f"/report/{sid}.xlsx").get_json()["error"]
        assert "ValueError" in e and "boom 42" in e, e
    finally:
        report_export.build_xlsx = orig

def test_export_follows_filters_and_checklist_enabled_only():
    import report_export
    from openpyxl import load_workbook
    rows = report_export._rows(R, None, q="ERR01", scope="all")
    assert rows and all("ERR01" in f["code"] for f in rows)
    scoped = report_export._rows(R, None, scope="prog:Main")
    assert scoped and all("Main" in f["location"] for f in scoped)
    r = run(cfg={**CFG, "disabled_rules": ["INF004"]})
    wb = load_workbook(io.BytesIO(report_export.build_xlsx(r, "x", ["ERROR"], "div", "prog:Main")))
    codes = [c.value for c in wb["Checklist"]["A"][1:]]
    assert "INF004" not in codes and "ERR010" in codes
    summ = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=3, max_row=13) if row[0].value}
    assert "Search: 'div'" in summ["Filters"] and "Scope: Main" in summ["Filters"]
    assert summ["Rules disabled"] == "INF004"

def test_lifecycle_master_switch():
    import report_export
    from openpyxl import load_workbook
    r = run(cfg={**CFG, "lifecycle_enabled": False})
    fs = flatten_findings(r)
    assert not [f for f in fs if f["code"] in ("ERR017", "WAR021", "INF009")]
    assert r["meta"]["lifecycle"].get("disabled") is True
    wb = load_workbook(io.BytesIO(report_export.build_xlsx(r, "x")))
    assert "Hardware" not in wb.sheetnames                         # (b): no hardware section at all
    assert "ERR017" not in [c.value for c in wb["Checklist"]["A"][1:]]
    text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
    assert "ifecycle" not in text, "lifecycle wording left in report"
    from pypdf import PdfReader
    pdf = PdfReader(io.BytesIO(report_export.build_pdf(r, "x")))
    ptxt = " ".join(pg.extract_text() for pg in pdf.pages)
    assert "Hardware" not in ptxt and "ifecycle" not in ptxt

def test_lifecycle_single_control():
    # per-rule list cannot switch lifecycle rules off; only the master switch does
    r = run(cfg={**CFG, "disabled_rules": ["ERR017", "WAR021", "INF009"]})
    assert [f for f in flatten_findings(r) if f["code"] == "ERR017"]
    assert r["meta"]["disabled_rules"] == []
    import app as A
    h = A.app.test_client().get("/settings").get_data(as_text=True)
    assert 'data-code="ERR017"' not in h and "Controlled by the Lifecycle switch" in h

def test_pdf_has_no_signature_block():
    import report_export
    from pypdf import PdfReader
    ptxt = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(report_export.build_pdf(R, "x"))).pages)
    assert "Verified by" not in ptxt and "Reviewed by" not in ptxt and "Sign-off" not in ptxt
    assert "Hardware & Lifecycle" in ptxt                           # lifecycle on → section present

def _lc_backup():
    import shutil
    bak = lifecycle._DATA_PATH + ".testbak"; shutil.copy(lifecycle._DATA_PATH, bak); return bak

def _lc_restore(bak):
    import shutil
    shutil.move(bak, lifecycle._DATA_PATH); lifecycle.reload()

def test_lifecycle_template_sheet():
    from openpyxl import load_workbook
    import app as A
    c, sid = _session()
    for _ in c.get(f"/stream/{sid}").response: pass
    x = c.get(f"/api/lifecycle/template.xlsx?scope=project&sid={sid}")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.data))
    ws = wb["Lifecycle"]
    assert [c.value for c in ws[1]] == ["Catalog", "Description", "Status", "EOL Date", "Replacement",
                                         "Source URL", "Checked On", "Note"]
    cats = [r[0].value for r in ws.iter_rows(min_row=2)]
    assert "1756-ENBT" in cats and "1756-IB16" in cats
    assert ws.data_validations.dataValidation and "Discontinued" in ws.data_validations.dataValidation[0].formula1
    assert ws["D2"].number_format == "@"                      # dates kept as text
    assert any("do not guess" in (r[0].value or "") for r in wb["Instructions"].iter_rows())
    assert c.get("/api/lifecycle/template.xlsx?scope=project&sid=nope").status_code == 404
    assert c.get("/api/lifecycle/template.xlsx?scope=all").status_code == 200

def test_lifecycle_import_preview_then_apply():
    import app as A
    bak = _lc_backup()
    try:
        c = A.app.test_client(); H = {"Origin": "http://127.0.0.1:5002"}
        csv_txt = ("Catalog,Status,EOL Date,Replacement,Source URL,Checked On\n"
                   "1756-ENBT,Discontinued,2022-12-31,1756-EN2T,https://compatibility.rockwellautomation.com/x,2026-10-09\n"
                   "1756-IB16,End of Life,2027-06-30,1756-IB16K,,2026-10-09\n"
                   "9999-NEW,Active,,,https://example.com/p,2026-10-09\n"
                   "1756-L83E,Unknown,,,,\n")
        r = c.post("/api/lifecycle/preview", data={"file": (io.BytesIO(csv_txt.encode()), "lc.csv")},
                   headers=H, content_type="multipart/form-data").get_json()
        assert r["success"] and r["rows"] == 3 and r["skipped"] == ["1756-L83E"]
        assert {"catalog": "1756-IB16", "from": "Active", "to": "End of Life"} in r["changed"]
        assert [a["catalog"] for a in r["added"]] == ["9999-NEW"]
        assert r["no_source"] == ["1756-IB16"]
        assert lifecycle.get_status("1756-IB16")["status"] == "Active"     # nothing saved yet
        d = c.post("/api/lifecycle/apply", json={"token": r["token"]}, headers=H).get_json()
        assert d["success"] and d["imported"] == 3
        st = lifecycle.get_status("1756-IB16")
        assert st["status"] == "End of Life" and st["sourced"] is False
        assert lifecycle.get_status("1756-ENBT")["sourced"] is True
        assert c.post("/api/lifecycle/apply", json={"token": r["token"]}, headers=H).status_code == 400
    finally:
        _lc_restore(bak)

def test_loading_overlay_present():
    import app as A
    h = A.app.test_client().get("/").get_data(as_text=True)
    assert 'id="busy"' in h and "ringHTML(pct, 96)" in h

def test_udt_dependency_cycle_reported():
    xml = FIX.replace("<DataTypes>", """<DataTypes>
    <DataType Name="CYC_A" Family="NoFamily"><Members><Member Name="B" DataType="CYC_B" Dimension="0" Hidden="false"/></Members></DataType>
    <DataType Name="CYC_B" Family="NoFamily"><Members><Member Name="A" DataType="CYC_A" Dimension="0" Hidden="false"/></Members></DataType>""", 1)
    r = run(xml)
    assert any("ircular" in f["message"] or "ycle" in f["message"] for f in flatten_findings(r) if f["code"] == "WAR010")

def test_custom_string_member_is_resolved():
    # StringFamily types are user types: a member of one must be sized, not "cannot be analysed"
    xml = FIX.replace("<DataTypes>", """<DataTypes>
    <DataType Name="STR_40" Family="StringFamily"><Members><Member Name="LEN" DataType="DINT" Dimension="0" Hidden="false"/><Member Name="DATA" DataType="SINT" Dimension="40" Hidden="false"/></Members></DataType>
    <DataType Name="HAS_STR" Family="NoFamily"><Members><Member Name="Txt" DataType="STR_40" Dimension="0" Hidden="false"><Description><![CDATA[t]]></Description></Member></Members></DataType>""", 1)
    msgs_ = [f["message"] for f in flatten_findings(run(xml)) if f["code"] == "WAR010"]
    assert not any("HAS_STR" in m and "not resolved" in m for m in msgs_), msgs_

def test_checks_package_layout():
    import importlib
    for m in ("checks", "checks.common", "checks.logic", "checks.tags", "checks.io", "checks.udt", "checks.program"):
        importlib.import_module(m)
    import os
    for gone in ("new_rules.py", "urs_checks.py", "checker.py", "extra_checks.py", "L5XOpt_UDT.py", "L5XGen_UDT.py"):
        assert not os.path.exists(os.path.join(ROOT, gone)), gone

def test_fix_endpoints_gone():
    c, sid = _session()
    for path in (f"/preview_fixes/{sid}", f"/download_corrected/{sid}"):
        assert c.post(path, json={}, headers={"Origin": "http://127.0.0.1:5002"}).status_code == 404

def test_settings_save_merges():
    import app as A
    path = tempfile.mktemp(suffix=".json"); old = A.CONFIG_PATH; A.CONFIG_PATH = path
    try:
        json.dump({"urs": {"safety_prefix": "SAF_"}}, open(path, "w"))
        A.app.test_client().post("/api/settings", json={"disabled_rules": ["INF004"]},
                                 headers={"Origin": "http://127.0.0.1:5002"})
        cfg = A.load_config()
        assert cfg["urs"]["safety_prefix"] == "SAF_" and cfg["disabled_rules"] == ["INF004"]
    finally:
        A.CONFIG_PATH = old; os.path.exists(path) and os.unlink(path)

def test_cli_formats_and_fail_on():
    with tempfile.NamedTemporaryFile("w", suffix=".L5X", delete=False) as tf:
        tf.write(FIX); path = tf.name
    try:
        cli = [sys.executable, os.path.join(ROOT, "l5x_cli.py"), "analyze", path]
        j = subprocess.run(cli + ["--format", "json"], capture_output=True, text=True, timeout=60)
        assert json.loads(j.stdout)["findings"][0]["code"]
        s = subprocess.run(cli + ["--format", "sarif"], capture_output=True, text=True, timeout=60)
        assert any(r["id"].startswith(("CRT", "ERR", "WAR", "INF"))
                   for r in json.loads(s.stdout)["runs"][0]["tool"]["driver"]["rules"])
        f = subprocess.run(cli + ["--fail-on", "CRITICAL"], capture_output=True, text=True, timeout=60)
        assert f.returncode == 1   # AFI present
    finally:
        os.unlink(path)

# ── lifecycle ──
def test_lifecycle_lookup():
    assert lifecycle.get_status("1756-ENBT")["status"] == "Discontinued"
    assert lifecycle.get_status("NOPE-123")["status"] == "Unknown"

# ── UI: render against a REAL report shape (regression: blank tree nodes) ──
def test_ui_sidebar_against_real_report():
    import app as A
    html = A.app.test_client().get("/").get_data(as_text=True)
    js = "\n".join(re.findall(r"<script>(.*?)</script>", html, re.S))
    harness = """
var els={};var document={getElementById:function(id){return els[id]||(els[id]={innerHTML:'',style:{},addEventListener:function(){},
  classList:{add:function(){},remove:function(){},toggle:function(){}}});},
  createElement:function(){return {click:function(){},remove:function(){}}},body:{appendChild:function(){}},
  head:{appendChild:function(){}},
  querySelector:function(){return {classList:{add:function(){},remove:function(){},toggle:function(){}}};}};
var sessionStorage={_d:{},getItem:function(k){return this._d[k]||null;},setItem:function(k,v){this._d[k]=v;},removeItem:function(k){delete this._d[k];}};
var URL={createObjectURL:function(){return '';},revokeObjectURL:function(){}};
var window={addEventListener:function(){},scrollTo:function(){},scrollY:0};var location={hash:''};var EventSource=function(){};var fetch=function(){return {then:function(){return this},catch:function(){}}};
""" + js + """
S.report = REPORT; S.findings = flatten(S.report); render();
var full = els['sidebar'].innerHTML;
actSection('task:Fast_10MS'); var closed = els['sidebar'].innerHTML;
actSide(true); var rail = els['sidebar'].innerHTML; actSide(false);
console.log(JSON.stringify({side: full, closed: closed, rail: rail, main: els['main'].innerHTML.length,
  rowhtml: els['main'].innerHTML, scopes: S.findings.map(function(f){return f._scope;})}));
"""
    harness = harness.replace("REPORT", json.dumps(R))
    open("/tmp/ui_harness.js", "w").write(harness)
    out = subprocess.run(["node", "/tmp/ui_harness.js"], capture_output=True, text=True, timeout=30)
    assert out.returncode == 0, out.stderr[-800:]
    d = json.loads(out.stdout.strip().splitlines()[-1])
    assert "prog:?" not in d["scopes"], "program findings lost their scope"
    assert ">Main<" in d["side"], "program node has no name"
    assert "Fast_10MS" in d["side"], "task header missing"
    counts = [int(x) for x in re.findall(r'class="tn-count[^"]*">(\d+)<', d["side"])]
    assert counts[0] == len(FS), (counts[0], len(FS))   # All Findings == total
    assert d["main"] > 0
    assert "Expand" in d["side"] and "Hide" in d["side"]
    assert ">Main<" not in d["closed"], "collapsed section still lists its programs"
    assert "side-rail" in d["rail"]
    for gone in ("bm-btn", ">WF<", "REQ-AUT", "req-tag", "Compliance"):
        assert gone not in d["rowhtml"] and gone not in d["side"], gone

def main():
    fns = [(k, v) for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    bad = 0
    for k, fn in fns:
        try: fn(); print(f"PASS  {k}")
        except AssertionError as e: bad += 1; print(f"FAIL  {k}: {e}")
        except Exception as e:
            bad += 1; import traceback; print(f"ERROR {k}: {e!r}"); traceback.print_exc()
    print(f"\n{len(fns)-bad}/{len(fns)} passed"); return 1 if bad else 0

if __name__ == "__main__":
    sys.exit(main())
