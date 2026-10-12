# lifecycle.py — module lifecycle subsystem.
"""
Lifecycle status for Rockwell catalog numbers, backed by a local JSON
dataset (`lifecycle.json`). The dataset is user-maintained: populate it
by importing a JSON/CSV built from Rockwell's Product Compatibility &
Download Center (PCDC) lifecycle pages, via the UI (Settings page) or
the CLI (`l5x_cli.py lifecycle-import`).

Design choice: no live scraping. PCDC requires authentication for most
lifecycle data, scraping breaks whenever Rockwell redesigns, and a lint
tool should not need network access. A curated local dataset that the
user refreshes deliberately is more reliable and auditable — the
`scraped_at` date is shown in the UI so staleness is visible.

Public API:
  is_available()              -> bool
  dataset_info()              -> dict
  get_status(catalog)         -> dict
  get_status_batch(catalogs)  -> list[dict]
  import_dataset(records, meta=None, merge=True) -> dict (result summary)
  parse_csv(text)             -> (records: dict, errors: list)
  reload()                    -> None
  lifecycle_findings(modules, sev) -> list[Finding]
"""

import os, json, csv, io, threading, datetime
from typing import Optional
from base_check import Finding, effective_severity

_DATA_PATH = os.path.join(os.path.dirname(__file__), "lifecycle.json")
_lock = threading.Lock()
_cache = {"loaded": False, "data": {}, "meta": {}}

VALID_STATUSES = {"Active", "Active Mature", "End of Life", "Discontinued"}

# CSV column aliases — PCDC exports and hand-built sheets vary.
_COL_ALIASES = {
    "catalog":     {"catalog", "catalog number", "catalognumber", "catalog_number", "part", "part number"},
    "status":      {"status", "lifecycle", "lifecycle status", "lifecycle_status", "life cycle status"},
    "eol_date":    {"eol_date", "eol", "eol date", "end of life date", "discontinued date", "end_date", "final ship date"},
    "replacement": {"replacement", "replacement catalog", "migration", "suggested replacement", "successor"},
    "source_url":  {"source url", "source_url", "source", "url", "reference"},
    "checked_on":  {"checked on", "checked_on", "checked", "verified on", "date checked"},
    "note":        {"note", "notes", "comment", "comments"},
}


# ─── Loading ───────────────────────────────────────────────────────────────

def _load():
    with _lock:
        if _cache["loaded"]:
            return
        _cache["loaded"] = True
        if not os.path.exists(_DATA_PATH):
            return
        try:
            with open(_DATA_PATH, "r", encoding="utf-8") as f:
                raw = json.load(f)
        except (OSError, json.JSONDecodeError):
            return
        _cache["meta"] = raw.pop("_meta", {})
        _cache["data"] = raw


def reload():
    """Force re-read from disk (after an import)."""
    with _lock:
        _cache["loaded"] = False
        _cache["data"] = {}
        _cache["meta"] = {}
    _load()


def is_available() -> bool:
    _load()
    return bool(_cache["data"])


def dataset_info() -> dict:
    _load()
    return {
        "available":   bool(_cache["data"]),
        "scraped_at":  _cache["meta"].get("scraped_at"),
        "count":       len(_cache["data"]),
        "source_note": _cache["meta"].get("source_note", ""),
        "verified":    _cache["meta"].get("verified", False),
    }


# ─── Lookup ────────────────────────────────────────────────────────────────

def get_status(catalog: str) -> dict:
    _load()
    catalog = (catalog or "").strip()
    if not catalog:
        return _unknown(catalog, "empty catalog")
    if not _cache["data"]:
        return _unknown(catalog, "lifecycle dataset not present — import one in Settings")
    entry = _cache["data"].get(catalog)
    if entry is None:
        # Try case-insensitive fallback
        cu = catalog.upper()
        for k, v in _cache["data"].items():
            if k.upper() == cu:
                entry = v
                break
    if entry is None:
        return _unknown(catalog, "no entry in lifecycle dataset")
    status = entry.get("status", "Unknown")
    return {
        "catalog":     catalog,
        "status":      status,
        "replacement": entry.get("replacement"),
        "eol_date":    entry.get("eol_date"),
        "source_url":  entry.get("source_url"),
        "checked_on":  entry.get("checked_on"),
        # A row is 'sourced' only if it cites a web page (e.g. Rockwell PCDC).
        "sourced":     str(entry.get("source_url") or "").lower().startswith("http"),
        "confidence":  "High" if status in VALID_STATUSES else "Medium",
        "note":        entry.get("note", ""),
    }


def _unknown(catalog: str, why: str) -> dict:
    return {
        "catalog":     catalog,
        "status":      "Unknown",
        "replacement": None,
        "eol_date":    None,
        "source_url":  None,
        "confidence":  "N/A",
        "note":        why,
    }


# ─── Lifecycle → findings ──────────────────────────────────────────────────

def lifecycle_findings(modules: list, sev: dict) -> list:
    """
    Turn module lifecycle statuses into analysis findings:
      Discontinued → module_lifecycle at config severity (default ERROR)
      End of Life  → one level softer (default WARNING)
      Active Mature→ INFO note (replacement planning)
    Unknown / Active produce nothing.
    """
    s = effective_severity("module_lifecycle", sev)
    if s == "DISABLED" or not is_available():
        return []

    softer = {"CRITICAL": "ERROR", "ERROR": "WARNING",
              "WARNING": "INFO", "INFO": "INFO"}.get(s, "WARNING")

    findings = []
    seen = set()
    for m in modules:
        cat = m.get("catalog") or ""
        if not cat or cat in seen:
            continue
        seen.add(cat)
        rec = get_status(cat)
        status = rec["status"]
        if status == "Discontinued":
            repl = f" Suggested replacement: {rec['replacement']}." if rec.get("replacement") else ""
            eol = f" (final availability {rec['eol_date']})" if rec.get("eol_date") else ""
            findings.append(Finding(
                "lc_discontinued", "Discontinued Module In Use", s,
                f"Module '{m.get('name','?')}' uses catalog {cat}, which is "
                f"discontinued{eol}.{repl} Spares may be unobtainable.",
                location=f"Module:{m.get('name','?')}"))
        elif status == "End of Life":
            repl = f" Suggested replacement: {rec['replacement']}." if rec.get("replacement") else ""
            eol = f" (EOL date {rec['eol_date']})" if rec.get("eol_date") else ""
            findings.append(Finding(
                "lc_eol", "End-of-Life Module In Use", softer,
                f"Module '{m.get('name','?')}' uses catalog {cat}, which has "
                f"reached end of life{eol}.{repl} Plan migration.",
                location=f"Module:{m.get('name','?')}"))
        elif status == "Active Mature":
            findings.append(Finding(
                "lc_mature", "Mature Module In Use", "INFO",
                f"Module '{m.get('name','?')}' uses catalog {cat} (Active Mature) "
                f"— newer alternatives exist; consider for new designs.",
                location=f"Module:{m.get('name','?')}"))
    return findings


# ─── Importing ─────────────────────────────────────────────────────────────

def _parse_table(header, rows):
    """
    Shared CSV/XLSX row parser. Returns (records, skipped, errors):
      records — {catalog: entry}
      skipped — catalogs left "Unknown" (not imported, never guessed)
      errors  — row-level problems (row skipped)
    """
    errors, skipped, records = [], [], {}
    head = [str(h or "").strip().lower() for h in header]
    colmap = {}
    for field, aliases in _COL_ALIASES.items():
        for i, h in enumerate(head):
            if h in aliases:
                colmap[field] = i
                break
    if "catalog" not in colmap:
        return {}, [], [f"no recognisable Catalog column in header: {list(header)}"]
    if "status" not in colmap:
        return {}, [], [f"no recognisable Status column in header: {list(header)}"]

    def get(row, field):
        i = colmap.get(field)
        if i is None or i >= len(row) or row[i] is None:
            return ""
        v = row[i]
        if hasattr(v, "strftime"):                 # Excel date cell
            return v.strftime("%Y-%m-%d")
        return str(v).strip()

    for n, row in enumerate(rows, start=2):
        if not any(str(c or "").strip() for c in row):
            continue
        cat = get(row, "catalog")
        if not cat:
            errors.append(f"row {n}: missing catalog")
            continue
        raw = get(row, "status")
        if raw.lower() in ("unknown", "", "n/a", "?"):
            skipped.append(cat)
            continue
        status = _normalise_status(raw)
        if status is None:
            errors.append(f"row {n} ({cat}): unrecognised status {raw!r}")
            continue
        records[cat] = {
            "status":      status,
            "eol_date":    get(row, "eol_date") or None,
            "replacement": get(row, "replacement") or None,
            "source_url":  get(row, "source_url") or None,
            "checked_on":  get(row, "checked_on") or None,
            "note":        get(row, "note") or "",
        }
    return records, skipped, errors


def parse_csv(text: str):
    """CSV → (records, errors). 'Unknown' rows are reported as errors-free skips."""
    try:
        rows = list(csv.reader(io.StringIO(text)))
    except csv.Error as e:
        return {}, [f"CSV parse error: {e}"]
    if not rows:
        return {}, ["empty file"]
    records, skipped, errors = _parse_table(rows[0], rows[1:])
    return records, errors


def parse_upload(filename: str, data: bytes):
    """xlsx or csv upload → (records, skipped, errors)."""
    fn = (filename or "").lower()
    if fn.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb["Lifecycle"] if "Lifecycle" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            return {}, [], ["empty sheet"]
        return _parse_table(rows[0], rows[1:])
    if fn.endswith(".csv"):
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        if not rows:
            return {}, [], ["empty file"]
        return _parse_table(rows[0], rows[1:])
    return {}, [], ["upload an .xlsx or .csv file"]


def preview_import(records: dict) -> dict:
    """Compare parsed records with the current dataset — nothing is saved."""
    _load()
    changed, added, unchanged = [], [], 0
    for cat, e in sorted(records.items()):
        cur = _cache["data"].get(cat)
        if cur is None:
            added.append({"catalog": cat, "to": e["status"]})
        elif (cur.get("status"), cur.get("eol_date"), cur.get("replacement")) != \
             (e["status"], e.get("eol_date"), e.get("replacement")):
            changed.append({"catalog": cat, "from": cur.get("status"), "to": e["status"]})
        else:
            unchanged += 1
    no_source = sorted(c for c, e in records.items()
                       if not str(e.get("source_url") or "").lower().startswith("http"))
    return {"changed": changed, "added": added, "unchanged": unchanged, "no_source": no_source}


def build_template_xlsx(catalogs: list, title: str) -> bytes:
    """
    Lifecycle sheet for an LLM / manual refresh. catalogs: [(catalog, description)].
    Pre-filled from the current dataset; Status is a dropdown; dates are text.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    _load()
    wb = Workbook()
    ws = wb.active
    ws.title = "Lifecycle"
    head = ["Catalog", "Description", "Status", "EOL Date", "Replacement",
            "Source URL", "Checked On", "Note"]
    for j, h in enumerate(head, 1):
        c = ws.cell(1, j, h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A5F")
    for i, (cat, desc) in enumerate(catalogs, 2):
        e = _cache["data"].get(cat) or {}
        vals = [cat, desc, e.get("status", "Unknown"), e.get("eol_date") or "",
                e.get("replacement") or "", e.get("source_url") or "",
                e.get("checked_on") or "", e.get("note") or ""]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name="Arial", size=10)
            c.number_format = "@"                          # text: no date/ID auto-conversion
    last = max(2, len(catalogs) + 1)
    dv = DataValidation(type="list", allow_blank=False,
                        formula1='"Active,Active Mature,End of Life,Discontinued,Unknown"')
    dv.error = "Choose one of: Active, Active Mature, End of Life, Discontinued, Unknown"
    ws.add_data_validation(dv)
    dv.add(f"C2:C{last}")
    for col, w in zip("ABCDEFGH", [20, 34, 16, 12, 18, 48, 12, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{last}"

    ins = wb.create_sheet("Instructions")
    ins.column_dimensions["A"].width = 110
    lines = [
        ("L5X Lint — lifecycle refresh", True),
        (title, False),
        ("", False),
        ("1. Fill Status, EOL Date (YYYY-MM-DD), Replacement, Source URL and Checked On for each row.", False),
        ("2. Every changed row needs a Source URL to a Rockwell page (PCDC or a product notice).", False),
        ("   Rows without a source are imported but marked unverified.", False),
        ("3. Leave Status = Unknown when not sure. Unknown rows are not imported.", False),
        ("4. Import back in L5X Lint: Settings → Lifecycle → Import. You review the changes before saving.", False),
        ("", False),
        ("Prompt to paste into an online assistant together with this file:", True),
        ("For each catalog number on the 'Lifecycle' sheet, find its current Rockwell Automation "
         "lifecycle status (Active, Active Mature, End of Life, or Discontinued) from Rockwell's "
         "Product Compatibility and Download Center or official product notices. Fill Status, "
         "EOL Date (YYYY-MM-DD, if announced), Replacement (Rockwell's suggested successor), "
         "Source URL (the exact page you used) and Checked On (today, YYYY-MM-DD). If you cannot "
         "find an official source, set Status to Unknown — do not guess. Return the completed "
         "sheet as .xlsx or .csv with the same columns.", False),
    ]
    for i, (txt, bold) in enumerate(lines, 1):
        c = ins.cell(i, 1, txt)
        c.font = Font(name="Arial", size=10, bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalise_status(raw: str) -> Optional[str]:
    s = (raw or "").strip().lower()
    if s in ("active", "preferred"):
        return "Active"
    if s in ("active mature", "mature"):
        return "Active Mature"
    if s in ("end of life", "eol", "end-of-life"):
        return "End of Life"
    if s in ("discontinued", "disc", "obsolete", "retired"):
        return "Discontinued"
    return None


def validate_records(records: dict):
    """Validate JSON-shaped records. Returns (clean, errors)."""
    clean, errors = {}, []
    if not isinstance(records, dict):
        return {}, ["records must be an object keyed by catalog number"]
    for cat, entry in records.items():
        if cat == "_meta":
            continue
        if not isinstance(entry, dict):
            errors.append(f"{cat}: entry must be an object")
            continue
        status = _normalise_status(entry.get("status", ""))
        if status is None:
            errors.append(f"{cat}: unrecognised status {entry.get('status')!r}")
            continue
        clean[cat] = {
            "status":      status,
            "eol_date":    entry.get("eol_date") or None,
            "replacement": entry.get("replacement") or None,
            "note":        entry.get("note", "") or "",
        }
        if entry.get("source_url"):
            clean[cat]["source_url"] = entry["source_url"]
        if entry.get("checked_on"):
            clean[cat]["checked_on"] = entry["checked_on"]
    return clean, errors


def import_dataset(records: dict, meta: dict = None, merge: bool = True) -> dict:
    """
    Write records into lifecycle.json (atomically). If merge=True, new
    records override existing ones but unrelated entries survive.
    Returns a summary {imported, total, merged, errors}.
    """
    clean, errors = validate_records(records)
    if not clean:
        return {"imported": 0, "total": 0, "merged": merge,
                "errors": errors or ["no valid records"]}

    _load()
    with _lock:
        existing = dict(_cache["data"]) if merge else {}
        existing.update(clean)

        new_meta = dict(_cache["meta"]) if merge else {}
        new_meta.update(meta or {})
        new_meta.setdefault("source_note",
            "User-imported lifecycle data. Verify against Rockwell PCDC.")
        new_meta["scraped_at"] = (meta or {}).get("scraped_at") or \
            datetime.date.today().isoformat()
        new_meta["schema_version"] = 1

        out = {"_meta": new_meta}
        out.update(dict(sorted(existing.items())))

        tmp = _DATA_PATH + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(out, f, indent=2)
        os.replace(tmp, _DATA_PATH)

        # refresh cache
        _cache["data"] = existing
        _cache["meta"] = new_meta
        _cache["loaded"] = True

    return {"imported": len(clean), "total": len(existing),
            "merged": merge, "errors": errors}
